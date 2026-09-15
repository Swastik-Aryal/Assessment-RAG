"""Load the Knowledge Base xlsx into retrieval-ready Chunk objects (one per Q&A row)."""
import logging
import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl

log = logging.getLogger(__name__)

TEXT_COLS = (
    "Section Heading",
    "Control Heading",
    "Question Text",
    "Answer",
    "Notes/Comment",
)
UUID_RE = re.compile(
    r"[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}"
)


@dataclass
class Chunk:
    """One KB row as a retrieval document. `id` is identifier or row-N."""
    id: str
    text: str
    original_id: str
    section: str
    control: str


def norm(v) -> str:
    """Strip, collapse whitespace, drop replacement chars. Whitespace-only becomes empty."""
    if v is None:
        return ""
    s = str(v).replace("\ufffd", "").replace("\xa0", " ").strip()
    return re.sub(r"\s+", " ", s)


def load_kb(path: Path) -> list[Chunk]:
    """Read the KB xlsx: find the Question Text/Answer header, skip empty questions."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    header_i, colmap = None, {}
    for i, row in enumerate(rows):
        cells = [norm(c) for c in row]
        if "Question Text" in cells and "Answer" in cells:
            header_i = i
            colmap = {name: j for j, name in enumerate(cells) if name}
            break
    if header_i is None:
        raise ValueError(f"no header row in {path}")

    chunks = []
    empty_run = 0
    for n, row in enumerate(rows[header_i + 1 :], start=header_i + 2):
        mapped = {name: norm(row[j]) if j < len(row) else "" for name, j in colmap.items()}
        if not any(mapped.values()):
            empty_run += 1
            if empty_run >= 20:
                # trailing blank block, not a sparse middle section
                break
            continue
        empty_run = 0
        q = mapped.get("Question Text", "")
        if not q:
            continue
        ident = mapped.get("identifier", "") or f"row-{n}"
        lines = [f"{h}: {mapped[h]}" for h in TEXT_COLS if mapped.get(h)]
        chunks.append(
            Chunk(
                ident,
                "\n".join(lines),
                mapped.get("Original ID", ""),
                mapped.get("Section Heading", ""),
                mapped.get("Control Heading", ""),
            )
        )
    log.info("loaded %s kb chunks from %s", len(chunks), path)
    return chunks


if __name__ == "__main__":
    import dataclasses, json
    from pipeline.config import load_settings

    chunks = load_kb(load_settings().KB_PATH)
    out = Path("outputs/kb_chunks.json")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps([dataclasses.asdict(c) for c in chunks], indent=2, ensure_ascii=False))
    print(f"{len(chunks)} chunks -> {out}")
