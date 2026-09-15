"""Extract questions from an Excel workbook using a WorkbookSchema."""
import logging

from openpyxl.utils import column_index_from_string

from pipeline.excel.render import cell_str
from pipeline.excel.schema import TableSchema, WorkbookSchema
from pipeline.models import Question

log = logging.getLogger(__name__)


def col_idx(letter: str) -> int:
    """Excel column letter to 1-based index (A=1)."""
    return column_index_from_string(letter)


def table_last_row(ws, table: TableSchema, stop_before: int | None = None) -> int:
    """Last non-empty row in this table's span. Always stop before the next table."""
    end = table.last_data_row or (ws.max_row or 1)
    if stop_before:
        end = min(end, stop_before - 1)
    c0, c1 = col_idx(table.first_col), col_idx(table.last_col)
    last = table.first_data_row
    for r in range(table.first_data_row, end + 1):
        if any(cell_str(ws.cell(r, c).value) for c in range(c0, c1 + 1)):
            last = r
    return last


def extract_table(ws, table: TableSchema, stop_before: int | None = None) -> list[Question]:
    """Pull questions from one table. Section banners update metadata only."""
    qcol = col_idx(table.question_col)
    header_q = cell_str(ws.cell(table.header_row, qcol).value)
    last = table_last_row(ws, table, stop_before)
    idcol = col_idx(table.id_col) if table.id_col else None
    seccol = col_idx(table.section_label_col) if table.section_label_col else None
    section = ""
    if seccol:
        for r in range(table.header_row + 1, table.first_data_row):
            lab = cell_str(ws.cell(r, seccol).value)
            if lab:
                section = lab
    out = []
    for r in range(table.first_data_row, last + 1):
        q = cell_str(ws.cell(r, qcol).value)
        if not q:
            if seccol:
                lab = cell_str(ws.cell(r, seccol).value)
                if lab:
                    section = lab
            continue
        if q == header_q:
            continue
        rid = cell_str(ws.cell(r, idcol).value) if idcol else ""
        out.append(
            Question(
                id=rid or f"{ws.title}!{r}",
                text=q,
                sheet=ws.title,
                row=r,
                section_label=section,
                table_id=table.table_id,
            )
        )
    return out


def extract(wb, schema: WorkbookSchema) -> list[Question]:
    """Extract questions from every table in the schema."""
    questions = []
    for sheet in schema.sheets:
        if sheet.sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet.sheet_name]
        tables = sheet.tables
        for i, table in enumerate(tables):
            nxt = tables[i + 1].header_row if i + 1 < len(tables) else None
            questions.extend(extract_table(ws, table, nxt))
    return questions


if __name__ == "__main__":
    import argparse
    import json
    from collections import defaultdict

    from openpyxl import load_workbook

    from pipeline.config.config import load_settings
    from pipeline.config.logging_setup import setup_logging
    from pipeline.excel.analyzer import infer_schema
    from pipeline.llm.gemini import GeminiClient

    p = argparse.ArgumentParser()
    p.add_argument("path")
    p.add_argument("--force-full", action="store_true")
    a = p.parse_args()
    setup_logging()
    s = load_settings()
    out = s.OUTPUT_DIR / "excel"
    schema = infer_schema(a.path, GeminiClient(s), out_dir=out, force_full=a.force_full)
    wb = load_workbook(a.path, data_only=False)
    try:
        qs = extract(wb, schema)
    finally:
        wb.close()
    by = defaultdict(int)
    for q in qs:
        by[q.sheet] += 1
        print(f"{q.sheet}\t{q.row}\t{q.id}\t{q.section_label[:40]}\t{q.text[:80]}")
    with open(out / "questions.json", "w", encoding="utf-8") as f:
        json.dump([q.__dict__ for q in qs], f, indent=2)
    print({"n": len(qs), "by_sheet": dict(by), "keys_unique": len({q.key() for q in qs}) == len(qs)})
