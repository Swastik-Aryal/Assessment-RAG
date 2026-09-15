"""Write RAG answers into a copy of the original workbook and append a summary sheet."""
import argparse
import json
import logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font

from pipeline.excel.extractor import col_idx, extract
from pipeline.excel.schema import TableSchema, WorkbookSchema
from pipeline.models import GeneratedAnswer, Question
from pipeline.rag.generator import cell_values

log = logging.getLogger(__name__)
SUMMARY = "SecurityPal Summary"


def _anchor(ws, row: int, col: int):
    """Writable cell at (row, col). For merged cells, the top-left if it is on this row. This handles merged cells."""
    cell = ws.cell(row, col)
    if not isinstance(cell, MergedCell):
        return cell
    for mr in ws.merged_cells.ranges:
        if cell.coordinate in mr:
            if mr.min_row == row:
                return ws.cell(mr.min_row, mr.min_col)
            log.warning("merged cell %s is not anchored on row %s", cell.coordinate, row)
            return None
    return None


def _write_cell(ws, row: int, col: str, value: str):
    """Write `value` into column `col` on `row`, wrapping text. Skip if the merge is off-row."""
    cell = _anchor(ws, row, col_idx(col))
    if cell is None:
        return
    cell.value = value
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def _table_for(schema: WorkbookSchema, q: Question) -> TableSchema | None:
    """Table that extracted this question."""
    for sheet in schema.sheets:
        if sheet.sheet_name != q.sheet:
            continue
        for t in sheet.tables:
            if t.table_id == q.table_id:
                return t
    return None


def _fallback() -> GeneratedAnswer:
    """Empty answer when generation fails; keeps the rest of the run going."""
    return GeneratedAnswer(
        ans_status="needs_review",
        confidence="low",
        answer="",
        sources=[],
    )


def write_workbook(
    src: Path,
    schema: WorkbookSchema,
    questions: list[Question],
    answers: dict[str, GeneratedAnswer],
    dest: Path,
    placeholder: bool = False,
) -> Path:
    """Copy src to dest, fill llm/constant targets, append the SecurityPal Summary sheet."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(src.read_bytes())
    wb = load_workbook(dest, data_only=False)
    by_sheet: dict[str, list[Question]] = {}
    for q in questions:
        by_sheet.setdefault(q.sheet, []).append(q)
    for sheet_name, qs in by_sheet.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for q in qs:
            table = _table_for(schema, q)
            if table is None:
                continue
            got = cell_values(answers[q.key()], table.fill_targets) if q.key() in answers else {}
            for t in table.fill_targets:
                if t.strategy == "skip" or t.col == table.question_col:
                    continue
                if t.strategy == "constant":
                    val = t.constant_value or ""
                elif placeholder and q.key() not in answers:
                    val = ""
                elif t.target_id in got:
                    val = got[t.target_id]
                else:
                    continue
                _write_cell(ws, q.row, t.col, val)

    if SUMMARY in wb.sheetnames:
        del wb[SUMMARY]
    sm = wb.create_sheet(SUMMARY)
    headers = ["Sheet", "Ref", "Question", "Status", "Confidence", "Answer", "Sources"]
    for i, h in enumerate(headers, 1):
        cell = sm.cell(1, i, h)
        cell.font = Font(bold=True)
    for i, q in enumerate(questions, 2):
        ans = answers.get(q.key())
        sm.cell(i, 1, q.sheet)
        sm.cell(i, 2, q.id)
        sm.cell(i, 3, q.text)
        sm.cell(i, 4, ans.ans_status if ans else ("unanswerable" if placeholder else ""))
        sm.cell(i, 5, ans.confidence if ans else "")
        sm.cell(i, 6, ans.answer if ans else "")
        sm.cell(i, 7, ", ".join(ans.sources) if ans else "")
        for c in range(1, 8):
            sm.cell(i, c).alignment = Alignment(wrap_text=True, vertical="top")
    sm.column_dimensions["C"].width = 60
    sm.column_dimensions["F"].width = 40
    wb.save(dest)
    wb.close()
    return dest


def _hit_dump(h) -> dict:
    """JSON-serializable retrieval hit for the run folder."""
    c = h.chunk
    return {
        "id": c.id,
        "original_id": c.original_id,
        "score": h.score,
        "section": c.section,
        "control": c.control,
        "text": c.text,
    }


if __name__ == "__main__":
    from datetime import datetime, timezone

    from pipeline.config.config import load_settings
    from pipeline.excel.analyzer import infer_schema
    from pipeline.llm.gemini import GeminiClient
    from pipeline.llm.ollama import OllamaClient
    from pipeline.config.logging_setup import setup_logging
    from pipeline.rag.generator import Generator, build_system, build_user
    from pipeline.rag.index import build_retriever

    p = argparse.ArgumentParser()
    p.add_argument("path")
    p.add_argument("--placeholder", action="store_true")
    p.add_argument("--out")
    a = p.parse_args()
    folder = Path(a.path)
    files = sorted(p for p in folder.glob("*.xlsx") if not p.name.startswith("~$"))
    settings = load_settings()
    gemini = GeminiClient(settings)
    retriever = None
    gen = None
    if not a.placeholder:
        retriever = build_retriever(settings)
        gen = Generator(OllamaClient(settings), retriever)
    try:
        for src in files:
            run_id = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "_excel_" + src.stem
            out = (Path(a.out) / run_id) if a.out else settings.OUTPUT_DIR / run_id
            out.mkdir(parents=True, exist_ok=True)
            setup_logging(settings.LOG_LEVEL, out / "run.log")
            dest = out / f"Completed - {src.name}"
            if dest.resolve() == src.resolve():
                raise SystemExit("refusing to overwrite the original workbook")

            schema = infer_schema(src, gemini, out_dir=out)
            wb = load_workbook(src, data_only=False)
            try:
                qs = extract(wb, schema)
            finally:
                wb.close()
            (out / "questions.json").write_text(
                json.dumps([q.__dict__ for q in qs], indent=2), encoding="utf-8"
            )
            log.info("extracted %s questions from %s", len(qs), src.name)
            print("extracted", len(qs), "questions from", src.name)

            answers: dict[str, GeneratedAnswer] = {}
            retrieval: dict[str, list] = {}
            prompts: dict[str, dict] = {}
            llm_out: dict[str, dict] = {}
            if a.placeholder:
                write_workbook(src, schema, qs, {}, dest, placeholder=True)
            else:
                for i, q in enumerate(qs, 1):
                    table = _table_for(schema, q)
                    targets = table.fill_targets if table else []
                    log.info("[%s/%s] %s %s", i, len(qs), q.id, q.text[:80])
                    print(f"[{i}/{len(qs)}] {q.id}", flush=True)
                    hits, system, user = [], "", ""
                    try:
                        ans, hits, system, user, dump = gen.answer(
                            q, targets, schema.instructions_text
                        )
                    except Exception:
                        log.exception("answer failed %s", q.id)
                        hits = retriever.retrieve(q.text)
                        system = build_system(targets, schema.instructions_text)
                        user = build_user(q, hits)
                        ans = _fallback()
                        dump = ans.model_dump()
                        dump.update(dump.pop("extra", {}))
                    answers[q.key()] = ans
                    llm_out[q.key()] = dump
                    retrieval[q.key()] = [_hit_dump(h) for h in hits]
                    prompts[q.key()] = {"system": system, "user": user}
                (out / "answers.json").write_text(
                    json.dumps(llm_out, indent=2),
                    encoding="utf-8",
                )
                (out / "retrieval.json").write_text(
                    json.dumps(retrieval, indent=2), encoding="utf-8"
                )
                (out / "prompts.json").write_text(
                    json.dumps(prompts, indent=2), encoding="utf-8"
                )
                write_workbook(src, schema, qs, answers, dest)

            counts = {}
            for ans in answers.values():
                counts[ans.ans_status] = counts.get(ans.ans_status, 0) + 1
            print("wrote", dest)
            print("counts", counts or "placeholder")
            print("run", out)
    finally:
        if retriever:
            retriever.close()
