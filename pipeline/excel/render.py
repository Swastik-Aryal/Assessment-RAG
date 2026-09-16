"""Render an Excel workbook to a compact or full text representation for LLM analysis."""
import argparse
import logging
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

COMPACT_ROWS = 50
COMPACT_CELL = 120


def cell_str(v) -> str:
    """Normalize a cell value to a single-line string (empty if blank)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ").replace("\ufffd", "")).strip()


def is_question_like(s: str) -> bool:
    """True if text looks like a question (ends with ? or is longer than 40 chars)."""
    s = s.strip()
    return bool(s) and (s.endswith("?") or len(s) > 40)


def fill_flag(cell) -> str:
    """Return a fill marker for the LLM render, or '' if the cell has no useful fill."""
    try:
        f = cell.fill
        if not f or f.fill_type in (None, "none"):
            return ""
        fg = f.fgColor
        if fg is None:
            return ""
        kind = getattr(fg, "type", None)
        if kind == "rgb":
            rgb = str(getattr(fg, "rgb", "") or "")
            if len(rgb) >= 6 and rgb[-6:] not in ("000000", "FFFFFF") and rgb not in ("00000000", "0"):
                return f"fill:{rgb[-6:]}"
        elif kind == "theme":
            return f"fill:theme{getattr(fg, 'theme', '?')}"
        elif kind == "indexed":
            return f"fill:idx{getattr(fg, 'indexed', '?')}"
    except Exception:
        return "fill:?"
    return ""


def _list_values(ws, formula1) -> list[str]:
    """Resolve a data-validation formula to the dropdown values (inline list or cell range)."""
    if not formula1:
        return []
    s = str(formula1).strip()
    if s.startswith('"') and s.endswith('"'):
        return [x.strip() for x in s[1:-1].split(",") if x.strip()]
    try:
        from openpyxl.utils import range_boundaries

        if "!" in s:
            sheet_name, addr = s.split("!", 1)
            sheet_name = sheet_name.strip("'")
            src = ws.parent[sheet_name]
        else:
            addr, src = s, ws
        addr = addr.replace("$", "")
        min_col, min_row, max_col, max_row = range_boundaries(addr)
        out = []
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                t = cell_str(src.cell(r, c).value)
                if t:
                    out.append(t)
        return out
    except Exception:
        return [s]


def _validations(ws) -> dict[str, list[str]]:
    """Map column letter -> allowed list values from sheet data validations."""
    out: dict[str, list[str]] = {}
    dvs = getattr(ws.data_validations, "dataValidation", None) or []
    for dv in dvs:
        if dv.type != "list":
            continue
        vals = _list_values(ws, dv.formula1)
        if not vals:
            continue
        for spec in str(dv.sqref).split():
            try:
                from openpyxl.utils import range_boundaries

                min_col, _, max_col, _ = range_boundaries(spec.replace("$", ""))
            except Exception:
                continue
            for c in range(min_col, max_col + 1):
                out[get_column_letter(c)] = vals
    return out


def _col_stats(ws) -> tuple[dict[str, dict], int]:
    """Per-column stats plus the last row that has any value."""
    stats = {}
    n = max(ws.max_row or 0, 1)
    last_ne = 0
    for c in range(1, (ws.max_column or 1) + 1):
        empty = qlike = total_len = 0
        for r in range(1, n + 1):
            t = cell_str(ws.cell(r, c).value)
            if not t:
                empty += 1
            else:
                last_ne = max(last_ne, r)
                total_len += len(t)
                qlike += int(is_question_like(t))
        letter = get_column_letter(c)
        stats[letter] = {
            "pct_question_like": round(100 * qlike / n, 1),
            "mean_length": round(total_len / n, 1),
            "pct_empty": round(100 * empty / n, 1),
        }
    return stats, last_ne


def render_wb(wb, full: bool = False) -> str:
    """Turn a workbook into text the analyzer LLM can read (compact or full)."""
    sheets = []
    for ws in wb.worksheets:
        stats, last_ne = _col_stats(ws)
        sheets.append((ws, stats, last_ne))
    wb_last = max((n for _, _, n in sheets), default=0)
    parts = [f"workbook: last_nonempty_row={wb_last}"]
    row_cap = None if full else COMPACT_ROWS
    cell_cap = None if full else COMPACT_CELL
    for ws, stats, last_ne in sheets:
        parts.append(f"# Sheet: {ws.title}")
        parts.append(f"dims: {ws.max_row}x{ws.max_column} last_nonempty_row={last_ne}")
        merged = [str(m) for m in ws.merged_cells.ranges]
        if merged:
            parts.append("merged: " + ", ".join(merged))
        vals = _validations(ws)
        if vals:
            parts.append(
                "validations: "
                + "; ".join(f"{c}={' | '.join(v)}" for c, v in vals.items())
            )
        parts.append(
            "col_stats: "
            + "; ".join(
                f"{c} q={s['pct_question_like']}% empty={s['pct_empty']}% mean={s['mean_length']}"
                for c, s in stats.items()
            )
        )
        last_r = ws.max_row if row_cap is None else min(ws.max_row, row_cap)
        for r in range(1, last_r + 1):
            cells = []
            bold_short = 0
            for c in range(1, (ws.max_column or 1) + 1):
                cell = ws.cell(r, c)
                t = cell_str(cell.value)
                if not t:
                    continue
                shown = t if cell_cap is None or len(t) <= cell_cap else t[:cell_cap] + "…"
                flags = []
                if cell.font and cell.font.bold:
                    flags.append("bold")
                    if len(t) <= 40:
                        bold_short += 1
                ff = fill_flag(cell)
                if ff:
                    flags.append(ff)
                flag = (" [" + "] [".join(flags) + "]") if flags else ""
                cells.append(f'{get_column_letter(c)}{r}="{shown}"{flag}')
            if not cells:
                continue
            # several short bold cells on one row is usually a header, not a question
            prefix = "[header-like] " if bold_short >= 3 else ""
            parts.append(prefix + " ".join(cells))
        parts.append("")
    return "\n".join(parts).rstrip() + "\n"


def render(path, full: bool = False) -> str:
    """Load `path` and render it. Compact: first 50 rows, cells cut at 120 chars."""
    wb = load_workbook(path, data_only=False)
    try:
        return render_wb(wb, full=full)
    finally:
        wb.close()


if __name__ == "__main__":
    from pipeline.config.config import load_settings
    from pipeline.config.logging_setup import setup_logging

    p = argparse.ArgumentParser()
    p.add_argument("path")
    p.add_argument("--full", action="store_true")
    a = p.parse_args()
    setup_logging()
    s = load_settings()
    text = render(a.path, full=a.full)
    out = s.OUTPUT_DIR / "excel"
    out.mkdir(parents=True, exist_ok=True)
    dest = out / "render.txt"
    dest.write_text(text, encoding="utf-8")
    print("wrote", dest)
    print("sheets", text.count("# Sheet:"), "chars", len(text), "full", a.full)
