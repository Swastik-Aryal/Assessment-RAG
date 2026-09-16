"""Infer workbook schema via Gemini with sheet/header validation and retry."""
import argparse
import json
import logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from pipeline.excel.render import cell_str, render
from pipeline.excel.schema import WorkbookAnalysis, WorkbookSchema
from pipeline.llm.structured import call_structured

log = logging.getLogger(__name__)

SYSTEM = """You describe the layout of a messy security-questionnaire Excel workbook for a RAG-based answering system.
The rendering lists every non-empty cell, merged ranges, dropdowns, and per-column stats.
Do not assume a fixed header row or column. Layouts vary: banners, grey section bands, inconsistent headers, stacked or side-by-side tables.

IMPORTANT: needs_full_render: Always set to true if any schema rendering has less no of rows than 'last_nonempty_row'. more_info_reason says what you think is missing. Still fill sheets as far as you can. RECOMMNEDED even if a single sheet is incomplete.

Identify:
- cover / instructions / response-key sheets (tables empty)
- every question table: header_row (some may not have header, in which case header_row=null), first_data_row(highest priority, make sure it is correct, can be 1 if header is not there), last_data_row or null, first_col, last_col as letters
- question_col: the column whose cells are actually questions (end with ? or long text). col_stats.pct_question_like is evidence. Never pick Domain/Category/Section label columns.
- id_col if a column is mostly short ids like 1.1.1; else null. Do not invent an id_regex.
- section_label_col if empty question rows carry a section banner; null if unsure
- fill_targets: every column a human respondent/security personnel would type into. No extra instruction
  on the sheet is required. You can only have one column with role=answer. Other fillable columns (applicability, comments,
  extra closed fields) are their own targets(set role = other if unsure).
- strategy=llm for (answer, applicability, comments, evidence, notes, extras, i.e. columns that can be filled with knowledge, or can be filled with citations and notes or extremely simple instruction following).
- strategy = skip only for totally irrelevant/unfillable columns likw owner which the rag might not be able to answer.
- For each fill_target you decide format, allowed_values, and rule.
  Enums are always a priority. Enforce as much as possible. ALWAYS use them to ADHERE TO global instructions.
  If there is an explicit global rule/format/enum for a column name, always use that. Never override it.
  Column names may contain rules/format instructions. If true, use them as enums/rules EXACTLY as written in the column name . THEY ALWAYS override the global format/rule eg: (Y/N MEANS Y OR N).
  Write that choice into `rule` so the answering model follows it.
  Use free_text and empty allowed_values only as a last resort when no other option is available.

- instructions_text: combine How to complete / instruction sheets (applies to all tables)
- response_key: value/meaning pairs from a key sheet or an in-sheet legend

Columns are letters (A, B, C). table_id must be unique. Sheets keep their real names.
Always consider the fact that single sheet can have multiple tables, and each table can have different columns.
"""


def analyze(gemini, rendering: str, full: bool = False, prev=None, error: str = "") -> WorkbookAnalysis:
    """Ask Gemini for layout. `full` means the rendering already has every row."""
    bits = []
    if full:
        bits.append("FULL RENDER — every row and cell is below. needs_full_render must be false.")
    if error:
        bits.append("Your previous schema failed validation. Fix only these issues and keep the rest:")
        bits.append(error)
        if prev is not None:
            bits.append("Previous schema:")
            bits.append(prev.model_dump_json(indent=2))
    bits.append(rendering)
    return call_structured(gemini, SYSTEM, "\n\n".join(bits), WorkbookAnalysis)


def schema_errors(path, schema: WorkbookSchema) -> list[str]:
    """Sheet names must match the workbook. Each fill_target header must sit at header_row/col."""
    wb = load_workbook(path, data_only=False)
    try:
        errs = []
        names = list(wb.sheetnames)
        want = [s.sheet_name for s in schema.sheets]
        extra = [n for n in want if n not in names]
        missing = [n for n in names if n not in want]
        if extra:
            errs.append(f"schema sheets not in workbook: {extra}; workbook has {names}")
        if missing:
            errs.append(f"workbook sheets missing from schema: {missing}")
        for s in schema.sheets:
            if s.sheet_name not in wb.sheetnames:
                continue
            ws = wb[s.sheet_name]
            for table in s.tables:
                for t in table.fill_targets:
                    if not t.col or not t.header:
                        continue
                    try:
                        c = column_index_from_string(t.col)
                    except ValueError:
                        errs.append(
                            f"{s.sheet_name} {table.table_id}: bad col {t.col!r}"
                        )
                        continue
                    cell = cell_str(ws.cell(table.header_row, c).value)
                    claimed = cell_str(t.header)
                    if claimed.lower() != cell.lower():
                        errs.append(
                            f"{s.sheet_name} {table.table_id} {t.col}{table.header_row}: "
                            f"schema header {claimed!r} vs cell {cell!r}"
                        )
        return errs
    finally:
        wb.close()


def _schema(analysis: WorkbookAnalysis) -> WorkbookSchema:
    return WorkbookSchema.model_validate(
        analysis.model_dump(exclude={"needs_full_render", "more_info_reason"})
    )


def infer_schema(path, gemini, out_dir: Path | None = None, force_full: bool = False) -> WorkbookSchema:
    """Compact render then LLM. On validation failure, up to 3 full-render retries. Writes schema.json."""
    attempts = []
    rendering = render(path, full=force_full)
    analysis = analyze(gemini, rendering, full=force_full)
    attempts.append({"full": force_full, **analysis.model_dump()})
    if analysis.needs_full_render and not force_full:
        log.info("llm asked for full render: %s", analysis.more_info_reason or "(no reason)")
        rendering = render(path, full=True)
        analysis = analyze(gemini, rendering, full=True)
        attempts.append({"full": True, **analysis.model_dump()})
    schema = _schema(analysis)
    errs = schema_errors(path, schema)

    # Checks for validation errors.
    rendering = render(path, full=True)
    for i in range(3):
        if not errs:
            break
        msg = "\n".join(errs)
        log.warning("schema failed validation (retry %s/3):\n%s", i + 1, msg)
        analysis = analyze(gemini, rendering, full=True, prev=schema, error=msg)
        attempts.append({"full": True, "validation_retry": i + 1, **analysis.model_dump()})
        schema = _schema(analysis)
        errs = schema_errors(path, schema)
    if out_dir:
        out_dir.mkdir(parents=True, exist_ok=True)
        (out_dir / "schema.json").write_text(
            json.dumps({"attempts": attempts, "schema": schema.model_dump()}, indent=2),
            encoding="utf-8",
        )
    if errs:
        raise RuntimeError("schema validation failed after 3 retries:\n" + "\n".join(errs))
    return schema


if __name__ == "__main__":
    from pipeline.config.config import load_settings
    from pipeline.llm.gemini import GeminiClient
    from pipeline.config.logging_setup import setup_logging

    p = argparse.ArgumentParser()
    p.add_argument("path")
    p.add_argument("--force-full", action="store_true")
    a = p.parse_args()
    setup_logging()
    s = load_settings()
    out = s.OUTPUT_DIR / "excel"
    schema = infer_schema(a.path, GeminiClient(s), out_dir=out, force_full=a.force_full)
    print(schema.model_dump_json(indent=2))
    print("wrote", out / "schema.json")
