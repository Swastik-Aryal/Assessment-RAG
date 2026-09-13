import json
import logging
import re
import time
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from pipeline.excel.analyzer import infer_schema
from pipeline.excel.extractor import extract
from pipeline.excel.writer import _fallback, _hit_dump, _table_for, write_workbook
from pipeline.config.logging_setup import add_run_log, remove_run_log
from pipeline.mail.client import Attachment
from pipeline.mail.watcher import parse_portal_request, ref
from pipeline.models import GeneratedAnswer, Question
from pipeline.portal.browser import PortalBrowser
from pipeline.rag.generator import _portal_target, build_system, build_user, portal_text

log = logging.getLogger(__name__)
_PW = re.compile(r"password is:\s*([A-Za-z0-9]+)", re.I)
PLACEHOLDER = "Needs review."


def seconds_until_rotation() -> float:
    return 3600 - (time.time() % 3600)


def _run_dir(session_dir: Path, scenario: str, msg) -> tuple[Path, list]:
    gen_id = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + f"_{scenario}_{msg.id}"
    out = session_dir / gen_id
    out.mkdir(parents=True, exist_ok=True)
    fhs = [add_run_log(out / "run.log")]
    (out / "request.json").write_text(
        json.dumps(
            {
                "id": msg.id,
                "from": msg.sender,
                "subject": msg.subject,
                "body": msg.body,
                "received_at": msg.received_at,
                "attachments": [a.filename for a in msg.attachments],
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    log.info("generation %s", gen_id)
    return out, fhs


def _counts(answers: dict[str, GeneratedAnswer]) -> dict[str, int]:
    c: dict[str, int] = {}
    for a in answers.values():
        c[a.ans_status] = c.get(a.ans_status, 0) + 1
    return c


def _write_answers(out: Path, llm_out, retrieval, prompts):
    (out / "answers.json").write_text(json.dumps(llm_out, indent=2), encoding="utf-8")
    (out / "retrieval.json").write_text(json.dumps(retrieval, indent=2), encoding="utf-8")
    (out / "prompts.json").write_text(json.dumps(prompts, indent=2), encoding="utf-8")


def answer_all(generator, questions: list[Question], get_targets, instructions="", on_progress=None):
    """Answer each question. Generation failure -> needs_review. on_progress every 10."""
    answers, llm_out, retrieval, prompts = {}, {}, {}, {}
    n = len(questions)
    for i, q in enumerate(questions, 1):
        targets = get_targets(q)
        log.debug("[%s/%s] %s %s", i, n, q.id, q.text[:80])
        hits, system, user, dump = [], "", "", {}
        try:
            ans, hits, system, user, dump = generator.answer(q, targets, instructions)
        except Exception:
            log.exception("answer failed %s", q.id)
            try:
                hits = generator.retriever.retrieve(q.text)
            except Exception:
                hits = []
            system = build_system(targets, instructions)
            user = build_user(q, hits)
            ans = _fallback()
            dump = ans.model_dump()
            dump.update(dump.pop("extra", {}) or {})
        answers[q.id] = ans
        llm_out[q.id] = dump
        retrieval[q.id] = [_hit_dump(h) for h in hits]
        prompts[q.id] = {"system": system, "user": user}
        if i == n or i % 10 == 0:
            log.info("answered %s/%s", i, n)
        if on_progress and i % 10 == 0:
            on_progress()
    return answers, llm_out, retrieval, prompts


def completion(mail, settings, msg, counts, extra="", attachments=()):
    """Send Completed: <subject> with status counts and [ref:]."""
    bits = ", ".join(f"{k}={v}" for k, v in sorted(counts.items())) or "none"
    subject = f"Completed: {msg.subject}" + (f" {extra}" if extra else "")
    body = f"Finished. Status counts: {bits}\n{ref(msg.id)}"
    return mail.send(settings.CLIENT_ADDRESS, subject, body, attachments)


def request_password(mail, settings) -> str:
    seen = {m["id"] for m in mail.list_inbox()}
    mail.send(
        settings.PORTAL_ACCESS_ADDRESS,
        "PORTAL ACCESS REQUEST",
        "Requesting the current portal password.",
    )
    log.info("requested portal password")
    reply = mail.wait_for(settings.PORTAL_ACCESS_ADDRESS, seen)
    m = _PW.search(reply.body or "")
    if not m:
        raise RuntimeError("password reply had no password")
    log.info("got portal password")
    return m.group(1)


def _login(browser, mail, settings, req, tries=3):
    for i in range(tries):
        pw = request_password(mail, settings)
        if browser.login(req.portal_url, req.username, pw) == "ok":
            log.info("portal login ok")
            return
        log.warning("portal login failed, try %s", i + 1)
    raise RuntimeError("portal login failed")


def run_excel(msg, settings, watcher, gemini, generator, session_dir: Path, on_progress=None):
    out, fhs = _run_dir(session_dir, "excel", msg)
    try:
        snap = watcher.sent_ack(msg.id)
        att = next(
            (a for a in msg.attachments if a.filename.lower().endswith((".xlsx", ".xlsm"))),
            None,
        )
        if att is None:
            raise RuntimeError("excel request has no workbook")
        src = out / att.filename
        src.write_bytes(att.content)
        log.info("saved attachment %s", att.filename)

        log.info("inferring excel schema")
        schema = infer_schema(src, gemini, out_dir=out)
        wb = load_workbook(src, data_only=False)
        try:
            qs = extract(wb, schema)
        finally:
            wb.close()
        (out / "questions.json").write_text(
            json.dumps([q.__dict__ for q in qs], indent=2), encoding="utf-8"
        )
        log.info("extracted %s questions", len(qs))

        def get_targets(q):
            t = _table_for(schema, q)
            return t.fill_targets if t else []

        answers, llm_out, retrieval, prompts = answer_all(
            generator, qs, get_targets, schema.instructions_text, on_progress
        )
        _write_answers(out, llm_out, retrieval, prompts)
        dest = out / f"Completed - {src.name}"
        log.info("writing workbook")
        write_workbook(src, schema, qs, answers, dest)
        counts = _counts(answers)

        if not watcher.ack_alive(snap):
            log.warning("mailbox reset mid-run, not sending %s", msg.id)
            return out
        watcher.mail.send(
            settings.CLIENT_ADDRESS,
            f"Completed: {msg.subject}",
            f"Please find the completed questionnaire attached.\n{ref(msg.id)}",
            attachments=[Attachment(dest.name, dest.read_bytes())],
        )
        log.info("sent workbook %s", dest.name)
        if watcher.ack_alive(snap):
            completion(watcher.mail, settings, msg, counts, extra="(summary)")
            log.info("sent completion for %s %s", msg.id, counts)
        log.info("excel done %s", msg.id)
        return out
    finally:
        for fh in fhs:
            remove_run_log(fh)


def _portal_payload(qs, answers) -> dict[str, str]:
    out = {}
    for q in qs:
        ans = answers.get(q.id)
        text = portal_text(ans).strip() if ans else ""
        out[q.id] = text or PLACEHOLDER
    return out


def run_portal(msg, settings, watcher, generator, session_dir: Path, on_progress=None):
    out, fhs = _run_dir(session_dir, "portal", msg)
    browser = None
    try:
        snap = watcher.sent_ack(msg.id)
        req = parse_portal_request(msg, settings)
        log.info("opening portal")
        browser = PortalBrowser(settings, out_dir=out)
        _login(browser, watcher.mail, settings, req)
        qs = browser.questions()
        (out / "questions.json").write_text(
            json.dumps([q.__dict__ for q in qs], indent=2), encoding="utf-8"
        )
        log.info("read %s portal questions", len(qs))
        browser.close()
        browser = None

        portal_t = _portal_target()
        answers, llm_out, retrieval, prompts = answer_all(
            generator, qs, lambda q: [portal_t], on_progress=on_progress
        )
        _write_answers(out, llm_out, retrieval, prompts)
        counts = _counts(answers)
        payload = _portal_payload(qs, answers)

        wait = seconds_until_rotation()
        if wait < 30:
            log.info("password rotates in %.0fs, waiting", wait + 1)
            time.sleep(wait + 1)

        log.info("submitting portal answers")
        browser = PortalBrowser(settings, out_dir=out)
        for _ in range(3):
            _login(browser, watcher.mail, settings, req)
            state = browser.submit(payload)
            if state == "incomplete":
                log.info("portal incomplete, filling blanks")
                for q in browser.questions():
                    if not (payload.get(q.id) or "").strip():
                        payload[q.id] = PLACEHOLDER
                state = browser.submit(payload)
            if state == "submitted":
                log.info("portal submitted")
                break
            if state == "expired":
                log.info("portal session expired, retrying")
                browser.new_session()
                continue
            raise RuntimeError(f"portal submit: {state}")
        else:
            raise RuntimeError("portal expired 3 times")

        if not watcher.ack_alive(snap):
            log.warning("mailbox reset mid-run, not sending %s", msg.id)
            return out
        completion(watcher.mail, settings, msg, counts)
        log.info("sent completion for %s %s", msg.id, counts)
        log.info("portal done %s", msg.id)
        return out
    finally:
        if browser:
            browser.close()
        for fh in fhs:
            remove_run_log(fh)
